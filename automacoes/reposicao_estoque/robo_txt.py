import openpyxl
import os
import glob
import sys
import re
import unicodedata
import requests
import traceback

# ==============================================================
# 🛡️ MÓDULO DE SEGURANÇA: VALIDAÇÃO GOOGLE SHEETS
# ==============================================================
def verificar_licenca(nome_robo):
    url_painel = "https://docs.google.com/spreadsheets/d/SEU_PAINEL_ID/pub?output=csv"
    
    print("==================================================")
    print(f"📡 Central da Diretoria: Validando credenciais de {nome_robo}...")
    
    try:
        resposta = requests.get(url_painel, timeout=12)
        resposta.raise_for_status() 
        texto_csv = resposta.text.upper()
        
        padrao_virgula = f"{nome_robo},ATIVO".upper()
        padrao_ponto_virgula = f"{nome_robo};ATIVO".upper()
        
        if (padrao_virgula in texto_csv) or (padrao_ponto_virgula in texto_csv):
            print(f"🔒 Licença validada: {nome_robo} autorizado via Control Room.\n")
            return True
        else:
            print("=" * 70)
            print(f"⛔ ACESSO NEGADO: {nome_robo} encontra-se DESATIVADO no painel principal.")
            print("=" * 70)
            os.system("pause")
            sys.exit(1)
            
    except requests.exceptions.RequestException as e:
        print("=" * 70)
        print(f"⛔ ERRO DE CONEXÃO: Falha ao contatar o servidor de licenciamento.")
        print(f"Detalhe técnico: {e}")
        print("=" * 70)
        os.system("pause")
        sys.exit(1)

# ==============================================================
# 🛠️ FUNÇÕES AUXILIARES DE TRATAMENTO
# ==============================================================
def remover_acentos(texto):
    if texto is None: return ""
    texto = str(texto).upper().strip()
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')

def limpar_referencia(val):
    if val is None: return ""
    if isinstance(val, (int, float)):
        return f"{val:.0f}"
    
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s

# ==============================================================
# 1. SETUP DE PASTAS
# ==============================================================
if getattr(sys, 'frozen', False):
    DIRETORIO_BASE = os.path.dirname(sys.executable)
else:
    DIRETORIO_BASE = os.path.dirname(os.path.abspath(__file__))

PASTA_LOJAS = os.path.join(DIRETORIO_BASE, "Lojas_Base")
PASTA_SAIDA = os.path.join(DIRETORIO_BASE, "TXTs_Exportados")

os.makedirs(PASTA_LOJAS, exist_ok=True)
os.makedirs(PASTA_SAIDA, exist_ok=True)

# ==============================================================
# 2. GERAÇÃO DOS TXTs (lê da ABA 2)
# ==============================================================
def gerar_txts():
    arquivos_brutos = glob.glob(os.path.join(PASTA_LOJAS, "*.xlsx")) + glob.glob(os.path.join(PASTA_LOJAS, "*.xlsm"))
    arquivos_loja = [f for f in arquivos_brutos if not os.path.basename(f).startswith('~$')]
    
    if not arquivos_loja:
        print(f"📭 Nenhum arquivo Excel encontrado na pasta '{PASTA_LOJAS}'.")
        return

    for caminho_loja in arquivos_loja:
        nome_arquivo = os.path.basename(caminho_loja)
        nome_sem_extensao = os.path.splitext(nome_arquivo)[0] 
        
        match = re.search(r'(H\d{3})', nome_arquivo.upper())
        if not match: continue
        
        codigo_h = match.group(1)
        print(f"📝 Gerando TXT para {nome_sem_extensao} ({codigo_h})...")
        
        wb = openpyxl.load_workbook(caminho_loja, data_only=True)
        
        if len(wb.worksheets) < 2:
            print(f"   ⚠️ Aba 2 não encontrada em {nome_sem_extensao}, pulando...")
            wb.close()
            continue
            
        ws2 = wb.worksheets[1]
        mapa2 = {}
        linha_cab2 = None
        
        for r in range(1, 15):
            for c in range(1, 15):
                val = remover_acentos(ws2.cell(row=r, column=c).value)
                if 'CODIGO' in val: mapa2['COD'] = c; linha_cab2 = r
                elif 'REPOSICAO' in val or 'REPO' in val: mapa2['REPO'] = c
            if linha_cab2: break
        
        linhas_txt = []
        
        if 'COD' in mapa2 and linha_cab2:
            for row in range(linha_cab2 + 1, ws2.max_row + 1):
                cod_celula_raw = ws2.cell(row=row, column=mapa2['COD']).value
                cod_val = limpar_referencia(cod_celula_raw)
                
                if not cod_val: continue
                
                rep_valor = 0
                if 'REPO' in mapa2:
                    rep_celula = ws2.cell(row=row, column=mapa2['REPO']).value
                    try:
                        rep_valor = int(float(str(rep_celula).replace(',', '.'))) if rep_celula is not None else 0
                    except ValueError:
                        rep_valor = 0
                
                if rep_valor > 0:
                    linhas_txt.append(f"{cod_val}998;{rep_valor}")

        wb.close()
        
        if linhas_txt:
            caminho_txt = os.path.join(PASTA_SAIDA, f"{nome_sem_extensao}.txt")
            with open(caminho_txt, 'w', encoding='utf-8') as f:
                f.write("\n".join(linhas_txt))
            print(f"   ✅ TXT gerado: {len(linhas_txt)} itens -> Salvo em TXTs_Exportados")
        else:
            print(f"   ⚠️ Nenhum item para reposição encontrado em {nome_sem_extensao}.")

    print("\n🏁 Geração de TXTs Finalizada!")
    print(f"📁 Arquivos salvos na pasta: {PASTA_SAIDA}")

# ==============================================================
# PONTO DE IGNIÇÃO COM TRAVA DE TELA
# ==============================================================
if __name__ == "__main__":
    try:
        verificar_licenca("TXT")
        gerar_txts()
    except Exception as e:
        print(f"\n❌ ERRO INESPERADO: {e}")
        traceback.print_exc()
    
    # Esta linha obriga a tela preta a ficar aberta no final
    print("\n")
    os.system("pause")