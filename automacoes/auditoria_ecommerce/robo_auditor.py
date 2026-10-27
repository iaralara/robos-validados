import os
import sys
import glob
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
import logging
import re

# =========================================================
# CONFIGURAÇÃO DE LOG E FUNÇÕES BLINDADAS
# =========================================================
logging.basicConfig(level=logging.INFO, format='%(message)s')
logger = logging.getLogger(__name__)

def converter_moeda_br(valor):
    if pd.isna(valor): return 0.0
    valor_str = str(valor).strip()
    valor_str = re.sub(r'[^\d,\.-]', '', valor_str)
    if ',' in valor_str and '.' in valor_str:
        valor_str = valor_str.replace('.', '').replace(',', '.')
    elif ',' in valor_str:
        valor_str = valor_str.replace(',', '.')
    try: return float(valor_str)
    except: return 0.0

def extrair_pedido(val):
    """Extrai o número do pedido sem falhar, mesmo se for float ou NaN."""
    if pd.isna(val): return '-'
    v = str(val).split('-')[0].strip()
    if v.endswith('.0'): v = v[:-2]  # Remove decimais fantasmas do Excel
    return v

def extrair_nf(val):
    if pd.isna(val): return '-'
    m = re.search(r'(\d+)', str(val))
    return m.group(1) if m else '-'

def obter_coluna(df, nome_col):
    """Garante que a coluna extraída nunca seja um DataFrame duplo, evitando erros."""
    if not nome_col or nome_col not in df.columns:
        return pd.Series(['-'] * len(df))
    col = df[nome_col]
    if isinstance(col, pd.DataFrame):
        return col.iloc[:, 0]
    return col

class MotorAuditoriaHibrido:
    def __init__(self):
        self.diretorio = os.path.dirname(os.path.abspath(__file__))
        self.arquivos_cigam = []
        self.arquivos_consolidado = []

    def identificar_arquivos(self):
        logger.info("=======================================")
        logger.info("🔍 MOTOR HÍBRIDO: RASTREAMENTO DE ARQUIVOS")
        logger.info("=======================================")
        
        todos_excel = glob.glob(os.path.join(self.diretorio, '*.xls*'))
        for arq in todos_excel:
            if 'Conciliacao_Processada' in arq or '~' in arq: continue
            
            nome_arquivo = os.path.basename(arq)
            nome_upper = nome_arquivo.upper()
            
            try:
                excel = pd.ExcelFile(arq)
                abas_limpas = [str(aba).strip().upper() for aba in excel.sheet_names]
                
                if 'RELATÓRIO' in abas_limpas or 'RELATORIO' in abas_limpas or 'CIGAM' in nome_upper:
                    self.arquivos_cigam.append(arq)
                    logger.info(f"   ✅ [CIGAM] Detectado: {nome_arquivo}")
                else:
                    self.arquivos_consolidado.append(arq)
                    logger.info(f"   ✅ [ALPARGATAS/OMNI] Detectado: {nome_arquivo}")
            except Exception as e:
                logger.warning(f"   ⚠️ Falha de leitura em {nome_arquivo}: {e}")

        if not self.arquivos_cigam or not self.arquivos_consolidado:
            logger.error("\n❌ ERRO CRÍTICO: Par de arquivos CIGAM e Consolidado não localizado.")
            sys.exit(1)

    def processar_dados(self):
        logger.info("\n🚀 Extraindo dados e blindando matrizes...")
        
        # ==========================================
        # 1. PROCESSAMENTO CIGAM
        # ==========================================
        lista_cigam = []
        for arq in self.arquivos_cigam:
            df = pd.read_excel(arq, sheet_name=0, skiprows=3)
            df.columns = [str(c).upper().strip() for c in df.columns]
            df = df.loc[:, ~df.columns.duplicated()].copy()
            lista_cigam.append(df)
        
        df_cigam_raw = pd.concat(lista_cigam, ignore_index=True)
        
        col_ped_cigam = next((c for c in df_cigam_raw.columns if 'COD PEDIDO' in c), None)
        df_cigam_raw['PEDIDO_CIGAM'] = obter_coluna(df_cigam_raw, col_ped_cigam).apply(extrair_pedido)

        col_nf = next((c for c in df_cigam_raw.columns if 'NF' in c or 'NOTA' in c), None)
        df_cigam_raw['NUM_NF'] = obter_coluna(df_cigam_raw, col_nf).apply(extrair_nf)

        col_emiss = next((c for c in df_cigam_raw.columns if 'EMISS' in c or 'DAT' in c), None)
        if col_emiss:
            df_cigam_raw['DATA_EMISSAO'] = pd.to_datetime(obter_coluna(df_cigam_raw, col_emiss), errors='coerce', dayfirst=True)
        else:
            df_cigam_raw['DATA_EMISSAO'] = pd.NaT

        col_frete_c = next((c for c in df_cigam_raw.columns if 'FRETE' in c), None)
        df_cigam_raw['FRETE_CIGAM'] = obter_coluna(df_cigam_raw, col_frete_c).apply(converter_moeda_br)

        col_tot_c = next((c for c in df_cigam_raw.columns if 'TOT ITEM' in c or 'VALOR' in c), None)
        df_cigam_raw['VALOR_CIGAM'] = obter_coluna(df_cigam_raw, col_tot_c).apply(converter_moeda_br)
        
        col_nome_c = next((c for c in df_cigam_raw.columns if 'NOME' in c or 'LOJA' in c), None)
        df_cigam_raw['LOJA_CIGAM'] = obter_coluna(df_cigam_raw, col_nome_c)

        def consolidar_nfs(x):
            validas = [str(i).strip() for i in x.unique() if pd.notna(i) and str(i).strip().upper() not in ['NAN', 'NONE', '-', '']]
            return ' / '.join(validas) if validas else '-'

        df_cigam = df_cigam_raw.groupby(['PEDIDO_CIGAM', 'LOJA_CIGAM'], dropna=False, as_index=False).agg({
            'DATA_EMISSAO': 'first',
            'VALOR_CIGAM': 'sum',
            'FRETE_CIGAM': 'sum',
            'NUM_NF': consolidar_nfs
        })

        # ==========================================
        # 2. PROCESSAMENTO OMNI E ALPARGATAS
        # ==========================================
        lista_alpa_pedidos = []
        
        for arq in self.arquivos_consolidado:
            dict_dfs = pd.read_excel(arq, sheet_name=None)
            abas_upper = {str(k).strip().upper(): k for k in dict_dfs.keys()}
            
            if 'BASE PEDIDOS' in abas_upper and 'BASE DATAS' in abas_upper:
                df_p = dict_dfs[abas_upper['BASE PEDIDOS']]
                df_d = dict_dfs[abas_upper['BASE DATAS']]
                
                df_p.columns = [str(c).upper().strip() for c in df_p.columns]
                df_d.columns = [str(c).upper().strip() for c in df_d.columns]
                
                df_p = df_p.loc[:, ~df_p.columns.duplicated()].copy()
                df_d = df_d.loc[:, ~df_d.columns.duplicated()].copy()
                
                col_ped = next((c for c in df_p.columns if 'NÚMERO DO PEDIDO' in c or 'PEDIDO' in c), None)
                if col_ped:
                    df_p['PEDIDO_ALPA'] = obter_coluna(df_p, col_ped).apply(extrair_pedido)
                
                if 'EMPRESA' in df_d.columns:
                    df_d = df_d[df_d['EMPRESA'].astype(str).str.upper() == 'ALPA']
                
                col_data = next((c for c in df_d.columns if 'DATA DO PAGAMENTO' in c or 'DATA PAGAMENTO' in c), None)
                col_mont = next((c for c in df_d.columns if 'MONT.MOEDA DOC' in c or 'MONTANTE' in c), None)
                
                df_d['DATA_PGTO'] = pd.to_datetime(obter_coluna(df_d, col_data), errors='coerce', dayfirst=True) if col_data else pd.NaT
                df_d['MONTANTE'] = obter_coluna(df_d, col_mont).apply(converter_moeda_br)
                    
                mapa_pgto = df_d.groupby(['FECHAMENTO', 'LOJA'], dropna=False, as_index=False).agg({'DATA_PGTO': 'first', 'MONTANTE': 'sum'})
                
                df_merged = pd.merge(df_p, mapa_pgto, on=['FECHAMENTO', 'LOJA'], how='left')
                lista_alpa_pedidos.append(df_merged)
                
            else:
                for aba, df in dict_dfs.items():
                    df_temp = pd.read_excel(arq, sheet_name=aba, header=None, nrows=20)
                    
                    has_pedido = df_temp.apply(lambda r: r.astype(str).str.contains(r'PEDIDO|DOCUMENTO', case=False, na=False).any(), axis=1).any()
                    if not has_pedido:
                        continue

                    mask = df_temp.apply(lambda r: r.astype(str).str.contains(r'PEDIDO|1\.FATURA|REPASSE|DOCUMENTO', case=False, na=False).any(), axis=1)
                    header_idx = df_temp[mask].index
                    
                    linha_cabecalho = header_idx[0] if not header_idx.empty else 0
                    df_aba = pd.read_excel(arq, sheet_name=aba, header=linha_cabecalho)
                    df_aba.columns = [str(c).upper().strip() for c in df_aba.columns]
                    
                    df_aba = df_aba.loc[:, ~df_aba.columns.duplicated()].copy()
                    
                    col_ped_aba = next((c for c in df_aba.columns if 'PEDIDO' in c or 'DOCUMENTO' in c), None)
                    if col_ped_aba:
                        df_aba['PEDIDO_ALPA'] = obter_coluna(df_aba, col_ped_aba).apply(extrair_pedido)
                        lista_alpa_pedidos.append(df_aba)

        # ==========================================
        # EXTRAÇÃO SEGURA DAS NOVAS COLUNAS
        # ==========================================
        alpa_master = []
        for df in lista_alpa_pedidos:
            res = pd.DataFrame()
            if 'PEDIDO_ALPA' not in df.columns: continue
            
            res['PEDIDO_ALPA'] = df['PEDIDO_ALPA']
            
            col_fech = next((c for c in df.columns if 'FECHAMENTO' in c), None)
            res['FECHAMENTO'] = obter_coluna(df, col_fech)
            
            col_loja = next((c for c in df.columns if 'LOJA' == c or 'NOME DA LOJA' in c), None)
            res['LOJA_ALPA'] = obter_coluna(df, col_loja)
            
            col_data_pg = next((c for c in df.columns if 'DATA_PGTO' == c or 'DATA PAGAMENTO' in c or 'DATA DO PAGAMENTO' in c), None)
            res['DATA_PGTO'] = pd.to_datetime(obter_coluna(df, col_data_pg), errors='coerce', dayfirst=True) if col_data_pg else pd.NaT
            
            col_montante = next((c for c in df.columns if 'MONTANTE' == c or 'MONT.MOEDA DOC' in c), None)
            res['MONTANTE'] = obter_coluna(df, col_montante).apply(converter_moeda_br)
            
            col_fat = next((c for c in df.columns if '1.FATURA' in c or ('FATURA' in c and 'NRO' not in c and 'STATUS' not in c)), None)
            res['FATURA'] = obter_coluna(df, col_fat).apply(converter_moeda_br)
            
            col_com = next((c for c in df.columns if 'COMISSAO' in c or 'COMISSÃO' in c or 'CAPTURA' in c), None)
            res['COMISSAO'] = obter_coluna(df, col_com).apply(converter_moeda_br)
            
            col_rep = next((c for c in df.columns if 'REPASSE' in c or 'MERCADORIA' in c or 'VALOR DO EVENTO' in c), None)
            res['REPASSE'] = obter_coluna(df, col_rep).apply(converter_moeda_br)
            
            col_fre = next((c for c in df.columns if 'FRETE' in c), None)
            res['FRETE_ALPA'] = obter_coluna(df, col_fre).apply(converter_moeda_br)
            
            alpa_master.append(res)
            
        df_alpa_full = pd.concat(alpa_master, ignore_index=True)
        
        df_alpa = df_alpa_full.groupby('PEDIDO_ALPA', dropna=False, as_index=False).agg({
            'FECHAMENTO': lambda x: next((i for i in x.dropna() if str(i).strip() != '-'), '-'),
            'LOJA_ALPA': lambda x: next((i for i in x.dropna() if str(i).strip() != '-'), '-'),
            'DATA_PGTO': 'first',
            'MONTANTE': 'sum',
            'FATURA': 'sum',
            'COMISSAO': 'sum',
            'REPASSE': 'sum',
            'FRETE_ALPA': 'sum'
        })

        # ==========================================
        # 3. CRUZAMENTO FINAL E PREPARAÇÃO
        # ==========================================
        df_final = pd.merge(df_cigam, df_alpa, left_on='PEDIDO_CIGAM', right_on='PEDIDO_ALPA', how='outer', indicator=True)

        df_final['PEDIDO'] = df_final['PEDIDO_CIGAM'].fillna(df_final['PEDIDO_ALPA'])
        df_final['LOJA'] = df_final['LOJA_CIGAM'].fillna(df_final['LOJA_ALPA']).fillna('-')
        
        df_final['VALOR_CIGAM'] = df_final['VALOR_CIGAM'].fillna(0)
        df_final['VALOR_DESC_15'] = df_final['VALOR_CIGAM'] * 0.85
        df_final['REPASSE'] = df_final['REPASSE'].fillna(0)
        df_final['DIFERENCA'] = df_final['REPASSE'] - df_final['VALOR_DESC_15']

        for col in ['DATA_EMISSAO', 'DATA_PGTO']:
            df_final[col] = pd.to_datetime(df_final[col], errors='coerce').dt.strftime('%d/%m/%Y').fillna('-')

        cols_conc = {
            'PEDIDO': 'Número do Pedido',
            'LOJA': 'Nome da Loja',
            'DATA_EMISSAO': 'Data Emissão CIGAM',
            'NUM_NF': 'Num NF',
            'FECHAMENTO': 'Fechamento',
            'VALOR_CIGAM': 'Valor CIGAM (R$)',
            'FRETE_CIGAM': 'Frete CIGAM (R$)',
            'VALOR_DESC_15': 'Valor c/ Desc. 15% (R$)',
            'DATA_PGTO': 'Data Pgto. Alpargatas',
            'MONTANTE': 'Montante (R$)',
            'FATURA': 'Fatura Alpargatas (R$)',
            'COMISSAO': 'Comissão Captura (R$)',
            'REPASSE': 'Repasse Mercadoria (R$)',
            'FRETE_ALPA': 'Frete Alpargatas (R$)',
            'DIFERENCA': 'Diferença (R$)'
        }

        df_conciliacao = df_final[df_final['_merge'] == 'both'].copy()
        df_so_cigam = df_final[df_final['_merge'] == 'left_only'].copy()
        df_so_alpa = df_final[df_final['_merge'] == 'right_only'].copy()

        return df_conciliacao[list(cols_conc.keys())].rename(columns=cols_conc), \
               df_so_cigam[['PEDIDO', 'LOJA', 'DATA_EMISSAO', 'NUM_NF', 'VALOR_CIGAM', 'FRETE_CIGAM', 'VALOR_DESC_15']].rename(columns={'PEDIDO': 'Número do Pedido', 'LOJA': 'Nome da Loja', 'DATA_EMISSAO': 'Data Emissão', 'NUM_NF': 'Num NF'}), \
               df_so_alpa[['PEDIDO_ALPA', 'LOJA', 'FECHAMENTO', 'DATA_PGTO', 'MONTANTE', 'FATURA', 'COMISSAO', 'REPASSE', 'FRETE_ALPA']].rename(columns={'PEDIDO_ALPA': 'Número do Pedido', 'LOJA': 'Nome da Loja', 'DATA_PGTO': 'Data Pgto.'})

    def gerar_excel(self, conc, cigam, alpa):
        agora = datetime.now().strftime("%d%m_%H%M%S")
        caminho = os.path.join(self.diretorio, f'Conciliacao_Processada_{agora}.xlsx')
        
        wb = Workbook()
        abas = [
            ("📊 Conciliação", conc, [6, 7, 8, 10, 11, 12, 13, 14, 15]),
            ("📊 Somente CIGAM", cigam, [5, 6, 7]),
            ("📊 Somente Alpargatas", alpa, [5, 6, 7, 8, 9])
        ]

        for i, (nome, df, col_moedas) in enumerate(abas):
            ws = wb.active if i == 0 else wb.create_sheet(nome)
            ws.title = nome
            
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)

            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal='center')
                
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for idx in col_moedas:
                    if idx <= len(row):  
                        row[idx-1].number_format = 'R$ #,##0.00'
                
                if nome == "📊 Conciliação":
                    dif = row[14] 
                    try:
                        if abs(float(dif.value or 0)) > 0.05:
                            dif.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                            dif.font = Font(color="9C0006", bold=True)
                    except: pass
                    
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 20
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

        wb.save(caminho)
        logger.info(f"✅ SUCESSO! Novo arquivo gerado e salvo em: {caminho}")

if __name__ == "__main__":
    motor = MotorAuditoriaHibrido()
    motor.identificar_arquivos()
    c, s_cigam, s_alpa = motor.processar_dados()
    motor.gerar_excel(c, s_cigam, s_alpa)