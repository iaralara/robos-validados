import gspread
from oauth2client.service_account import ServiceAccountCredentials
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from datetime import datetime, timedelta
import time
import re
import os
import sys
import traceback
import glob
import shutil
from collections import defaultdict
from dotenv import load_dotenv

# ============================================================================
# FUNCOES DE CAMINHO PARA O EXECUTAVEL ACHAR OS ARQUIVOS
# ============================================================================
def obter_caminho_arquivo(nome_arquivo):
    if getattr(sys, 'frozen', False):
        diretorio = os.path.dirname(sys.executable)
    else:
        diretorio = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(diretorio, nome_arquivo)

# Carrega o .env usando o caminho correto
load_dotenv(obter_caminho_arquivo(".env"))

# ============================================================================
# SENSOR DE FERIADOS
# ============================================================================
try:
    import holidays
    feriados_br = holidays.BR()
except ImportError:
    feriados_br = {
        datetime(2026, 1, 1).date(): "Ano Novo",
        datetime(2026, 2, 16).date(): "Carnaval",
        datetime(2026, 2, 17).date(): "Carnaval",
        datetime(2026, 4, 3).date(): "Paixão de Cristo",
        datetime(2026, 4, 21).date(): "Tiradentes",
        datetime(2026, 5, 1).date(): "Dia do Trabalho",
        datetime(2026, 6, 4).date(): "Corpus Christi",
        datetime(2026, 9, 7).date(): "Independência do Brasil",
        datetime(2026, 10, 12).date(): "Nossa Senhora Aparecida",
        datetime(2026, 11, 2).date(): "Finados",
        datetime(2026, 11, 15).date(): "Proclamação da República",
        datetime(2026, 12, 25).date(): "Natal"
    }

def eh_feriado(data_obj):
    return data_obj.date() in feriados_br

# ============================================================================
# CONFIGURACOES
# ============================================================================
print("=" * 70)
print("ROBO CIGAM - EXTRATOR DE PAGAMENTOS (TODAS LOJAS - 1 DOWNLOAD)".center(70))
print("=" * 70)

CIGAM_ACESSO = os.getenv("CIGAM_ACESSO")
CIGAM_USER = os.getenv("CIGAM_USER")
CIGAM_SENHA = os.getenv("CIGAM_SENHA")

if not all([CIGAM_ACESSO, CIGAM_USER, CIGAM_SENHA]):
    print("[ERRO FATAL] Credenciais ausentes. Verifique o arquivo .env.")
    input("Pressione ENTER para fechar...")
    sys.exit()

MESES_PT = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
    5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
    9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"
}

LOJAS_ALVO = [
    {"uf": "PB", "codigo_cigam": "L001", "nome_planilha": "MATRIZ - L001", "abas_destino": ["MES_ATUAL"]},
    {"uf": "PB", "codigo_cigam": "L002", "nome_planilha": "FILIAL 1 - L002", "abas_destino": ["MES_ATUAL"]},
    {"uf": "PE", "codigo_cigam": "L003", "nome_planilha": "FILIAL 2 - L003", "abas_destino": ["MES_ATUAL"]},
    {"uf": "AL", "codigo_cigam": "L004", "nome_planilha": "FILIAL 3 - L004", "abas_destino": ["MES_ATUAL"]},
    {"uf": "RN", "codigo_cigam": "L005", "nome_planilha": "FILIAL 4 - L005", "abas_destino": ["MES_ATUAL"]}
]

MAPA_LOJAS = {l["codigo_cigam"]: l for l in LOJAS_ALVO}

hoje = datetime.now()
if hoje.weekday() == 0:  
    print(f"\n[DATA] Hoje é {hoje.strftime('%A')}! Processando Sexta, Sábado e Domingo...")
    datas_processar = [hoje - timedelta(days=3), hoje - timedelta(days=2), hoje - timedelta(days=1)]
else:
    print(f"\n[DATA] Hoje é {hoje.strftime('%A')}. Processando apenas ontem...")
    datas_processar = [hoje - timedelta(days=1)]

print(f"[INFO] Datas a processar: {[d.strftime('%d/%m/%Y') for d in datas_processar]}")

PASTA_DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")

def autenticar_google():
    print("[LOGIN] Conectando ao Google Sheets...")
    escopo = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
    caminho_chave = obter_caminho_arquivo("chave_planilha.json")
    creds = ServiceAccountCredentials.from_json_keyfile_name(caminho_chave, escopo)
    return gspread.authorize(creds)

def clique_forcado(driver, xpath, descricao="elemento"):
    try:
        elemento = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH, xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
        time.sleep(0.3)
        driver.execute_script("arguments[0].click();", elemento)
        return True
    except:
        return False

# ============================================================================
# FUNCAO: Baixar e ler o arquivo Excel (1x para TODAS as lojas)
# ============================================================================

def baixar_e_ler_excel_geral(navegador, aguardar, data_str):
    resultados = {}
    
    arquivos_antes = set(glob.glob(os.path.join(PASTA_DOWNLOADS, "*.xlsx")))
    print(f"   [ARQUIVO] Arquivos .xlsx antes: {len(arquivos_antes)}")
    
    print("   [ACAO] Clicando em 'Salvar Arquivo'...", end=" ", flush=True)
    
    xpaths_salvar = [
        "//button[contains(text(), 'Salvar Arquivo')]",
        "//*[contains(text(), 'Salvar Arquivo')]",
        "//button[contains(@class, 'salvar')]",
        "//a[contains(text(), 'Salvar')]",
    ]
    
    clicou = False
    for xpath in xpaths_salvar:
        if clique_forcado(navegador, xpath, "Salvar Arquivo"):
            clicou = True
            break
    
    if not clicou:
        print("[ERRO] Nao conseguiu clicar em Salvar Arquivo")
        return resultados
    
    print("[OK]")
    
    print("   [AGUARDE] Aguardando download...", end=" ", flush=True)
    arquivo_baixado = None
    
    for i in range(30):
        time.sleep(2)
        
        arquivos_agora = set(glob.glob(os.path.join(PASTA_DOWNLOADS, "*.xlsx")))
        novos_arquivos = arquivos_agora - arquivos_antes
        
        if novos_arquivos:
            arquivo_baixado = max(novos_arquivos, key=os.path.getctime)
            try:
                tamanho1 = os.path.getsize(arquivo_baixado)
                time.sleep(1)
                tamanho2 = os.path.getsize(arquivo_baixado)
                if tamanho1 == tamanho2 and tamanho1 > 1000:
                    print(f"[OK]\n   [ARQUIVO] Arquivo: {os.path.basename(arquivo_baixado)} ({tamanho1} bytes)")
                    break
            except:
                pass
    
    if not arquivo_baixado:
        todos = glob.glob(os.path.join(PASTA_DOWNLOADS, "*.xlsx"))
        if todos:
            arquivo_baixado = max(todos, key=os.path.getctime)
            print(f"[AVISO]\n   [ARQUIVO] Usando arquivo mais recente: {os.path.basename(arquivo_baixado)}")
        else:
            print("[ERRO] Nenhum arquivo .xlsx encontrado!")
            return resultados
    
    print("   [LENDO] Lendo Excel...", end=" ", flush=True)
    
    try:
        import openpyxl
        wb = openpyxl.load_workbook(arquivo_baixado, data_only=True)
        ws = wb.active
        
        print(f"[OK] ({ws.max_row} linhas x {ws.max_column} cols)")
        
        loja_atual = None
        tem_dinheiro_na_loja = False
        
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            col_a = str(row[0]) if row[0] is not None else ""
            col_b = str(row[1]) if len(row) > 1 and row[1] is not None else ""
            col_c = row[2] if len(row) > 2 else None
            
            match_loja = re.search(r'Loja:\s*([A-Z]\d{3})', col_a, re.IGNORECASE)
            if match_loja:
                if loja_atual and loja_atual not in resultados:
                    resultados[loja_atual] = {"valor": 0.0, "apareceu": True, "tem_dinheiro": False}
                
                loja_atual = match_loja.group(1).upper()
                tem_dinheiro_na_loja = False
                continue
            
            if loja_atual and col_b.strip().lower() == "dinheiro":
                tem_dinheiro_na_loja = True
                if col_c is not None:
                    try:
                        if isinstance(col_c, (int, float)):
                            valor = float(col_c)
                        else:
                            v_str = str(col_c).strip().replace(',', '.')
                            valor = float(v_str)
                        
                        resultados[loja_atual] = {"valor": valor, "apareceu": True, "tem_dinheiro": True}
                        print(f"\n   [OK] {loja_atual}: R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
                        loja_atual = None 
                    except Exception as e:
                        print(f"\n   [AVISO] Erro ao converter valor de {loja_atual}: {col_c} - {e}")
                        resultados[loja_atual] = {"valor": 0.0, "apareceu": True, "tem_dinheiro": False}
                        loja_atual = None
        
        if loja_atual and loja_atual not in resultados:
            resultados[loja_atual] = {"valor": 0.0, "apareceu": True, "tem_dinheiro": False}
        
        wb.close()
        print(f"\n   [INFO] Total de lojas no relatorio: {len(resultados)}")
        
        pasta_processados = os.path.join(PASTA_DOWNLOADS, "processados_cigam")
        os.makedirs(pasta_processados, exist_ok=True)
        novo_nome = os.path.join(pasta_processados, f"GERAL_{data_str.replace('/','-')}_{os.path.basename(arquivo_baixado)}")
        shutil.move(arquivo_baixado, novo_nome)
        print(f"   [ARQUIVO] Arquivo movido para: {novo_nome}")
        
    except ImportError:
        print("[ERRO] openpyxl nao instalado. Instale com: pip install openpyxl")
    except Exception as e:
        print(f"[ERRO] Erro ao ler Excel: {e}")
        traceback.print_exc()
    
    return resultados

# ============================================================================
# FUNCAO: Obter valor de uma celula especifica na planilha (para rolagem)
# ============================================================================

def obter_valor_celula(matriz, linha, coluna):
    """Obtem valor numerico de uma celula da matriz, retornando 0.0 se vazia/invalida"""
    try:
        val_str = str(matriz[linha][coluna]).strip()
        if not val_str or "NENHUM" in val_str.upper() or "FERIADO" in val_str.upper():
            return 0.0
        texto_limpo = re.sub(r'[^\d,]', '', val_str)
        if not texto_limpo:
            return 0.0
        return float(texto_limpo.replace(',', '.'))
    except:
        return 0.0

# ============================================================================
# FUNCAO OTIMIZADA: Salvar no Google com cache e batch update
# ============================================================================

CACHE_PLANILHAS = {}

def salvar_no_google_otimizado(cliente, resultados_por_loja, data_alvo):
    dia_alvo = data_alvo.day
    nome_mes_alvo = MESES_PT[data_alvo.month]
    
    abas_para_atualizar = defaultdict(list)
    
    for nome_loja, resultado in resultados_por_loja.items():
        loja_info = None
        for l in LOJAS_ALVO:
            if l["nome_planilha"] == nome_loja:
                loja_info = l
                break
        
        if not loja_info:
            continue
        
        valor_cigam_num = resultado["valor"]
        apareceu = resultado["apareceu"]
        tem_dinheiro = resultado["tem_dinheiro"]
        
        valor_final_numerico = valor_cigam_num
        
        # === ROLAGEM: Soma valor do dia anterior se estiver entre 0.99 e 149.99 ===
        pode_rolar = True
        lojas_isentas = ["GUEDES", "PATOS", "SHOP SUL", "FC PB"]
        if any(isenta in nome_loja for isenta in lojas_isentas):
            pode_rolar = False
        if data_alvo.weekday() in [4, 5, 6]:  # Sexta, Sabado, Domingo nao rola
            pode_rolar = False
        
        if pode_rolar and valor_cigam_num > 0:
            try:
                aba_principal = loja_info["abas_destino"][0]
                nome_aba_rolagem = nome_mes_alvo if aba_principal == "MÊS_ATUAL" else aba_principal
                
                # Carrega cache da aba se nao estiver carregado
                cache_key = nome_aba_rolagem
                if cache_key not in CACHE_PLANILHAS:
                    planilha = cliente.open("Pagamentos").worksheet(nome_aba_rolagem)
                    CACHE_PLANILHAS[cache_key] = planilha.get_all_values()
                    time.sleep(0.5)
                
                matriz = CACHE_PLANILHAS[cache_key]
                valor_ontem_num = 0.0
                
                if nome_aba_rolagem == nome_mes_alvo:
                    # Abas mensais: procura linha da loja e coluna do dia anterior
                    linha_loja, col_dia = -1, -1
                    for r_idx, row in enumerate(matriz):
                        for c_idx, val in enumerate(row):
                            v_str = str(val).strip().upper()
                            if nome_loja.upper() in v_str:
                                linha_loja = r_idx
                            if v_str == str(dia_alvo) or v_str.startswith(f"{dia_alvo}/"):
                                col_dia = c_idx
                    
                    if linha_loja != -1 and col_dia > 0:
                        # Valor do dia anterior esta na coluna anterior (col_dia - 1)
                        valor_ontem_num = obter_valor_celula(matriz, linha_loja, col_dia - 1)
                else:
                    # Abas especificas: procura coluna da loja e linha do dia
                    col_loja = -1
                    for r_idx, row in enumerate(matriz):
                        for c_idx, val in enumerate(row):
                            if nome_loja.upper() in str(val).strip().upper():
                                col_loja = c_idx
                                break
                        if col_loja != -1:
                            break
                    
                    if col_loja != -1:
                        linha_dia = -1
                        for r_idx in range(len(matriz)):
                            v_str = str(matriz[r_idx][col_loja]).strip()
                            if v_str == str(dia_alvo) or v_str.startswith(f"{dia_alvo}/"):
                                linha_dia = r_idx
                                break
                        if linha_dia > 0:
                            # Valor do dia anterior esta na linha anterior, coluna ao lado
                            valor_ontem_num = obter_valor_celula(matriz, linha_dia - 1, col_loja + 1)
                
                if 0.99 <= valor_ontem_num <= 149.99:
                    valor_final_numerico += valor_ontem_num
                    c_str = f"{valor_cigam_num:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                    o_str = f"{valor_ontem_num:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                    t_str = f"{valor_final_numerico:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                    print(f"\n   [ROLAGEM] {nome_loja} | CIGAM (R$ {c_str}) + Ant (R$ {o_str}) = TOTAL: R$ {t_str}")
            
            except Exception as e:
                print(f"\n   [AVISO] ROLAGEM {nome_loja}: {str(e)[:60]}")
        
        # === FORMATA VALOR FINAL ===
        if not apareceu:
            valor_str_final = "Nenhum registro"
        elif not tem_dinheiro and valor_final_numerico == 0:
            valor_str_final = "0,00"
        else:
            valor_str_final = f"{valor_final_numerico:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        
        # === SALVA EM TODAS AS ABAS DESTINO ===
        for aba_original in loja_info["abas_destino"]:
            # CORRECAO: Define aba_principal ANTES de usar, nao depende do bloco de rolagem
            aba_principal = loja_info["abas_destino"][0]
            nome_aba = nome_mes_alvo if aba_original == "MÊS_ATUAL" else aba_principal
            
            cache_key = nome_aba
            if cache_key not in CACHE_PLANILHAS:
                try:
                    planilha = cliente.open("Pagamentos").worksheet(nome_aba)
                    CACHE_PLANILHAS[cache_key] = planilha.get_all_values()
                    time.sleep(0.5)
                except Exception as e:
                    print(f"\n   [ERRO] Erro ao carregar aba '{nome_aba}': {e}")
                    continue
            
            matriz = CACHE_PLANILHAS[cache_key]
            linha_alvo, coluna_alvo = -1, -1
            
            if nome_aba == nome_mes_alvo:
                # Abas mensais: loja na coluna A, dias na primeira linha
                for r_idx, row in enumerate(matriz):
                    for c_idx, val in enumerate(row):
                        v_str = str(val).strip().upper()
                        if nome_loja.upper() in v_str:
                            linha_alvo = r_idx + 1
                        if v_str == str(dia_alvo) or v_str.startswith(f"{dia_alvo}/"):
                            coluna_alvo = c_idx + 1
                
                if linha_alvo != -1 and coluna_alvo != -1:
                    abas_para_atualizar[nome_aba].append((linha_alvo, coluna_alvo, valor_str_final))
                else:
                    print(f"\n   [ERRO MAPEAMENTO] {nome_loja} nao encontrado na aba {nome_aba} (linha={linha_alvo}, col={coluna_alvo})")
            else:
                # Abas especificas: loja no cabecalho, dias na coluna da loja
                col_loja = -1
                for r_idx, row in enumerate(matriz):
                    for c_idx, val in enumerate(row):
                        if nome_loja.upper() in str(val).strip().upper():
                            col_loja = c_idx + 1
                            break
                    if col_loja != -1:
                        break
                
                if col_loja != -1:
                    linha_dia = -1
                    for r_idx in range(len(matriz)):
                        v_str = str(matriz[r_idx][col_loja - 1]).strip()
                        if v_str == str(dia_alvo) or v_str.startswith(f"{dia_alvo}/"):
                            linha_dia = r_idx + 1
                            break
                    if linha_dia != -1:
                        abas_para_atualizar[nome_aba].append((linha_dia, col_loja + 1, valor_str_final))
                    else:
                        print(f"\n   [ERRO MAPEAMENTO] Dia {dia_alvo} nao encontrado para {nome_loja} na aba {nome_aba}")
                else:
                    print(f"\n   [ERRO MAPEAMENTO] {nome_loja} nao encontrado no cabecalho de {nome_aba}")
    
    total_celulas = sum(len(v) for v in abas_para_atualizar.values())
    print(f"\n   [SALVANDO] Salvando {total_celulas} celulas em {len(abas_para_atualizar)} abas...")
    
    for nome_aba, celulas in abas_para_atualizar.items():
        try:
            planilha = cliente.open("Pagamentos").worksheet(nome_aba)
            
            batch_size = 50
            for i in range(0, len(celulas), batch_size):
                batch = celulas[i:i+batch_size]
                
                cell_list = []
                for linha, coluna, valor in batch:
                    cell = planilha.cell(linha, coluna)
                    cell.value = valor
                    cell_list.append(cell)
                
                planilha.update_cells(cell_list)
                time.sleep(1)
                
                for linha, coluna, valor in batch:
                    print(f"      [OK] {nome_aba}!{linha},{coluna} = {valor}")
            
        except Exception as e:
            print(f"\n   [ERRO] Erro ao salvar na aba '{nome_aba}': {str(e)[:80]}")
            time.sleep(5)

# ============================================================================
# FUNCAO: Processar TODAS as lojas em UM UNICO relatorio
# ============================================================================

def processar_todas_lojas(navegador, aguardar, data_alvo, cliente_google, somas_por_dia):
    data_formatada = data_alvo.strftime("%d/%m/%Y")
    
    print(f"\n{'='*70}")
    print(f"DATA: {data_formatada} | Processando TODAS as lojas de uma vez")
    print(f"{'='*70}")
    
    if eh_feriado(data_alvo):
        print(f"[FERIADO] {data_formatada} e FERIADO! Marcando todas as lojas...")
        resultados_feriado = {l["nome_planilha"]: {"valor": 0.0, "apareceu": False, "tem_dinheiro": False} for l in LOJAS_ALVO}
        salvar_no_google_otimizado(cliente_google, resultados_feriado, data_alvo)
        return True
    
    tentativa = 1
    max_tentativas = 5
    
    while tentativa <= max_tentativas:
        print(f"\n[TENTATIVA] Tentativa {tentativa}...", end=" ", flush=True)
        
        try:
            navegador.get("https://app.cigamgestor.com.br/")
            time.sleep(2)
            
            if not clique_forcado(navegador, "//a[contains(@class, 'menu-toggle')]", "menu"): 
                raise Exception("Falha no Menu")
            time.sleep(1)
            if not clique_forcado(navegador, "//*[contains(text(), 'Relat')]", "relatorios"): 
                raise Exception("Falha em Relatorios")
            time.sleep(1)
            if not clique_forcado(navegador, "//*[contains(text(), 'Vendas')]", "vendas"): 
                raise Exception("Falha em Vendas")
            time.sleep(1)
            if not clique_forcado(navegador, "//*[contains(text(), 'Formas')]", "formas"): 
                raise Exception("Falha em Formas")
            time.sleep(2)
            
            for id_campo in ["filtro_1", "filtro_2"]:
                campo = aguardar.until(EC.presence_of_element_located((By.ID, id_campo)))
                navegador.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
                time.sleep(0.3)
                navegador.execute_script("arguments[0].click();", campo)
                time.sleep(0.3)
                campo.send_keys(Keys.CONTROL + "a")
                campo.send_keys(Keys.BACKSPACE)
                time.sleep(0.2)
                for char in data_formatada:
                    campo.send_keys(char)
                    time.sleep(0.05)
                campo.send_keys(Keys.TAB)
                time.sleep(0.5)
            
            print("[ACAO] Visualizando relatorio GERAL...", end=" ", flush=True)
            if not clique_forcado(navegador, "//*[contains(text(), 'Visualizar')]", "visualizar"): 
                raise Exception("Botao Visualizar falhou")
            print("[OK]")
            
            print("[AGUARDE] Aguardando relatorio carregar...", end=" ", flush=True)
            time.sleep(5)
            
            for i in range(10):
                time.sleep(2)
                navegador.switch_to.default_content()
                
                try:
                    botoes = navegador.find_elements(By.XPATH, "//*[contains(text(), 'Salvar Arquivo')]")
                    if botoes:
                        print("[OK]")
                        break
                except:
                    pass
                
                try:
                    alertas = navegador.find_elements(By.XPATH, "//*[contains(translate(text(), 'NENHUM', 'nenhum'), 'nenhum')]")
                    if alertas:
                        print("[VAZIO] Vazio!")
                        resultados_vazio = {l["nome_planilha"]: {"valor": 0.0, "apareceu": False, "tem_dinheiro": False} for l in LOJAS_ALVO}
                        salvar_no_google_otimizado(cliente_google, resultados_vazio, data_alvo)
                        return True
                except:
                    pass
            else:
                print("[AVISO] Timeout, mas tentando baixar mesmo assim...")
            
            resultados_cigam = baixar_e_ler_excel_geral(navegador, aguardar, data_formatada)
            
            resultados_para_salvar = {}
            
            print(f"\n   [OK] Lojas no relatorio CIGAM: {len(resultados_cigam)}")
            print(f"   [INFO] Lojas a processar: {len(LOJAS_ALVO)}")
            
            for loja in LOJAS_ALVO:
                codigo = loja["codigo_cigam"]
                nome = loja["nome_planilha"]
                
                if codigo in resultados_cigam:
                    res = resultados_cigam[codigo]
                    valor = res["valor"]
                    
                    if res["tem_dinheiro"]:
                        v_str = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        print(f"   [VALOR] {codigo} ({nome}): R$ {v_str}")
                    else:
                        print(f"   [ZERO] {codigo} ({nome}): Apareceu mas sem Dinheiro -> 0,00")
                    
                    resultados_para_salvar[nome] = res
                    somas_por_dia[data_alvo] += valor
                else:
                    print(f"   [AUSENTE] {codigo} ({nome}): Nao apareceu no relatorio -> Nenhum registro")
                    resultados_para_salvar[nome] = {"valor": 0.0, "apareceu": False, "tem_dinheiro": False}
            
            salvar_no_google_otimizado(cliente_google, resultados_para_salvar, data_alvo)
            
            return True
            
        except Exception as e:
            print(f"\n[ERRO] Erro: {str(e)[:80]}")
            tentativa += 1
            time.sleep(3)
            continue
    
    print(f"\n[ERRO FATAL] ESGOTADAS {max_tentativas} TENTATIVAS")
    
    resultados_erro = {l["nome_planilha"]: {"valor": 0.0, "apareceu": False, "tem_dinheiro": False} for l in LOJAS_ALVO}
    try:
        salvar_no_google_otimizado(cliente_google, resultados_erro, data_alvo)
    except:
        pass
    
    return False

# ============================================================================
# EXECUCAO PRINCIPAL
# ============================================================================

navegador = None
cliente_google = None
contador_sucesso = 0
somas_por_dia = {data: 0.0 for data in datas_processar}
total_lojas = len(LOJAS_ALVO)

try:
    cliente_google = autenticar_google()
    print("[OK] Google Sheets conectado e pronto!")

    print("\n" + "=" * 70)
    print("[INICIANDO] Iniciando navegador...".center(70))
    print("=" * 70)
    
    opcoes = Options()
    opcoes.add_experimental_option("detach", True)
    opcoes.add_argument("--start-maximized")
    opcoes.add_argument("--no-sandbox")
    opcoes.add_argument("--disable-dev-shm-usage")
    opcoes.add_argument("--disable-blink-features=AutomationControlled")
    opcoes.add_experimental_option("excludeSwitches", ["enable-automation"])
    
    prefs = {
        "download.default_directory": PASTA_DOWNLOADS,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "safebrowsing.enabled": True
    }
    opcoes.add_experimental_option("prefs", prefs)
    
    navegador = webdriver.Chrome(options=opcoes)
    aguardar = WebDriverWait(navegador, 20)
    print("[OK] Navegador iniciado com sucesso!")
    
    print("\n" + "=" * 70)
    print("[LOGIN] Realizando login...".center(70))
    print("=" * 70)
    
    navegador.get("https://app.cigamgestor.com.br/login")
    time.sleep(3)
    
    try:
        aguardar.until(EC.presence_of_element_located((By.ID, "accesskey"))).send_keys(CIGAM_ACESSO)
        navegador.find_element(By.ID, "user").send_keys(CIGAM_USER)
        navegador.find_element(By.ID, "password").send_keys(CIGAM_SENHA)
        clique_forcado(navegador, "//*[@id='buttonLogin']", "botao login")
    except Exception as e:
        raise Exception("Verifique as credenciais no .env e a pagina de login.")
    
    time.sleep(5)
    print("[OK] Login realizado com sucesso!")
    
    print("\n" + "=" * 70)
    print(f"[INFO] Processando {total_lojas} lojas em 1 relatorio geral...".center(70))
    print("=" * 70)
    
    for data_alvo in datas_processar:
        CACHE_PLANILHAS.clear()
        
        sucesso = processar_todas_lojas(navegador, aguardar, data_alvo, cliente_google, somas_por_dia)
        if sucesso:
            contador_sucesso += total_lojas
        
        time.sleep(3)

except Exception as e:
    print(f"\n\n[ERRO CRITICO GLOBAL]: {e}")
    traceback.print_exc()

finally:
    print("\n" + "=" * 70)
    print("FINALIZANDO".center(70))
    print("=" * 70)
    print(f"[OK] Lojas processadas com sucesso: {contador_sucesso}/{total_lojas * len(datas_processar)}")
    
    print("\n" + "=" * 70)
    print("RESUMO DE VALORES EXTRAIDOS (CONFERENCIA)".center(70))
    print("=" * 70)
    
    soma_total_geral = 0.0
    for data_alvo in datas_processar:
        total_do_dia = somas_por_dia[data_alvo]
        soma_total_geral += total_do_dia
        total_formatado = f"{total_do_dia:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        print(f" DATA: {data_alvo.strftime('%d/%m/%Y')}: R$ {total_formatado}")
        
    print("-" * 70)
    total_geral_formatado = f"{soma_total_geral:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    print(f" SOMA TOTAL GERAL: R$ {total_geral_formatado}")
    print("=" * 70)
    
    if navegador:
        try:
            navegador.quit()
        except: pass
    print("\n[FIM] Robo finalizado!")
    
    input("\nPressione ENTER para fechar a janela...")